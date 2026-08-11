#!/usr/bin/env python3
"""Generate the WSQ "AI Vibe Coding for Multi-Agents System" (TGS-2020503207) Learner Guide as BOTH
a Markdown mirror (LG-*.md at repo root) and a DOCX (courseware/LG-*.docx) from one source.

House format: cover page, Document Version Control Record, auto TOC, Arial 11pt body, one section
per lab (Objective · Goal · What you'll build · DETAILED step-by-step with commands · Test it),
plus setup, revision and glossary.

NOTE ON DIVISION OF CONTENT: the slide deck is deliberately visual and carries NO step-by-step
walkthroughs. This Learner Guide is the single place where the detailed step-by-step instructions
for every lab live.
"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5
import prodoc
def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")

# ---------------- block DSL (single content stream → MD + DOCX) ----------------
B=[]
def h1(t): B.append(("h1",t))
def h2(t): B.append(("h2",t))
def h3(t): B.append(("h3",t))
def p(t):  B.append(("p",t))
def bullets(xs): B.append(("bullets",xs))
def steps(xs): B.append(("steps",xs))
def code(t): B.append(("code",t))
def note(t): B.append(("note",t))
def rule(): B.append(("rule",))

NLABS=len(ACT)

# ---------------- content ----------------
h1("Introduction")
p(f"This Learner Guide accompanies the WSQ course {C.TITLE} ({C.COURSE_CODE}), conducted by {C.ORG}. "
  f"It provides the detailed step-by-step instructions for all {NLABS} hands-on labs, organised by the "
  "five topics of the published course outline. Every lab maps to a learning outcome and is completed "
  "in Python on your own laptop.")
p("Use this guide alongside the course slides. The slide deck is deliberately visual — it shows the "
  "architecture, the workflow and what each lab produces — while this guide is the single place where "
  "the full step-by-step instructions and commands are recorded. Work through the steps here as the "
  "trainer walks through the concepts on the slides.")
p(f"All lab code is available in the course repository at {C.REPO_URL}. Keep this guide open during "
  "the assessment: the final assessment is open book.")

h1("Course Learning Outcomes")
bullets(C.LEARNING_OUTCOMES)

h1("Skills Framework Alignment")
p(f"TSC Title: {C.TSC_TITLE}    |    TSC Code: {C.TSC_CODE}")
bullets(C.TSC_ABILITIES)

h1("Before You Start — Environment Setup")
h3("What you need")
bullets([
 "A laptop (Windows, macOS or Linux) with Python 3.11 or later installed.",
 "A terminal and a code editor — VS Code is recommended.",
 "An OpenAI API key from platform.openai.com (used in Topics 1, 3 and 5).",
 "A Google AI Studio API key from aistudio.google.com (used in Topic 4).",
 "An AI coding agent for the vibe coding labs — Claude Code, Gemini CLI, Codex CLI, Cursor or Windsurf.",
 "Git, for version control during the vibe coding labs.",
 "A stable internet connection — every lab calls a hosted model API.",
])
h3("Create your working environment")
p("Create one project folder for the whole course with a single virtual environment. Every lab builds "
  "on the same environment, so you install the dependencies once.")
code("# create the course folder and a virtual environment\n"
     "mkdir multi-agent-course && cd multi-agent-course\n"
     "python3 -m venv .venv\n"
     "source .venv/bin/activate        # Windows: .venv\\Scripts\\activate\n\n"
     "# install everything the labs need\n"
     "pip install openai openai-agents google-adk pydantic python-dotenv streamlit \"mcp[cli]\"")
h3("Store your API keys safely")
p("Never hard-code an API key in source code, and never commit one to Git. Put your keys in a .env "
  "file and load them with python-dotenv. Add .env to .gitignore before your first commit.")
code("# .env  (never commit this file)\n"
     "OPENAI_API_KEY=sk-...\n"
     "GOOGLE_API_KEY=...\n\n"
     "# .gitignore\n"
     ".env\n.venv/\n__pycache__/\nmemory.json")
h3("Verify the setup")
p("Confirm Python and the key packages are importable before the first lab. If any import fails, "
  "re-activate the virtual environment and re-run the install command.")
code("python3 --version\n"
     "python3 -c \"import openai, agents, pydantic, streamlit; print('OK')\"\n"
     "python3 -c \"from google.adk.agents import Agent; print('ADK OK')\"\n"
     "python3 -c \"from mcp.server.fastmcp import FastMCP; print('MCP OK')\"")
h3("How the labs are organised")
p("Every lab has its own folder under labs/ in the course repository, named lab-01-…, lab-02-… and so "
  "on. Each folder contains a runnable Python script, a README with the lab instructions, and a Jupyter "
  "notebook you can open directly in Google Colab if you would rather not install anything locally.")
bullets([
 "Run locally: activate the virtual environment, ensure your .env holds the API keys, then run the "
 "script named in that lab's README (for example: python agent.py).",
 "Run in Colab: open the lab's README on GitHub and click the Open in Colab badge. Put your API keys "
 "in Colab Secrets rather than typing them into a cell.",
 "The two vibe coding labs (Labs 3 and 4) have no script to run — you direct an AI coding agent to "
 "write the code, so their folders hold the specification and context files instead.",
])
h3("Conventions used in every lab")
bullets([
 "Commands are run from your terminal with the virtual environment activated.",
 "Placeholders such as <YOUR_KEY> and /absolute/path/ are replaced with your own values.",
 "Model names change over time — if a model identifier is rejected, check the provider's current model list.",
 "Cost control: every lab uses small prompts, but keep an eye on your provider usage dashboard.",
])

# ---------------- per-topic, per-lab ----------------
for t in C.TOPICS:
    h1(f"Topic {t['code']} — {t['title']}")
    p(t["subtitle"])
    h3("Key concepts")
    bullets([f"{name} — {desc}" for name,desc in t["concepts"]])
    for a in [x for x in ACT if x["topic"]==t["num"]]:
        h2(f"Lab {a['num']} — {a['title']}")
        p(f"Learning outcome: {a['objective']}.")
        p(f"Goal: {a['desc']}")
        h3("What you'll build")
        p(a["build"]+f"   (Tools: {a['services']}.)")
        slug=C.lab_slug(a["num"], a["title"])
        h3("Where the code lives")
        bullets([
            f"Lab folder: labs/{slug}/ — the runnable Python script, a README and the notebook.",
            f"Run it locally: activate your virtual environment, set your keys in .env, then run the "
            f"script named in the lab's README.",
            f"Or open it in Google Colab: {C.colab_url(a['num'], a['title'])}",
        ])
        h3("Step-by-step")
        steps([(instr,cmd) for instr,cmd in a["steps"]])
        h3("Test it")
        p(a["test"])
        note(f"The complete working code for this lab is in labs/{slug}/ in the course repository. "
             f"Keep your API keys in .env — never commit them.")
        rule()

h1("Course Revision")
p("Use this checklist to revise before the assessment. Each point maps to a learning outcome and is "
  "covered by both the slides and the labs.")
bullets([
 "Describe the components of a modern AI agent and explain the reason-act-observe loop.",
 "Explain what agent skills and agent memory are, and when each is used.",
 "Explain what MCP is, what an MCP server exposes, and why the standard matters.",
 "State when a single agent is sufficient and when to split into a multi-agent system.",
 "List the stages of the structured vibe coding workflow and explain context engineering.",
 "Build an OpenAI Agents SDK agent with a function tool and a Pydantic output type.",
 "Build a triage agent that routes to specialist agents via handoffs.",
 "Build the equivalent coordinator and sub-agents in the Google Gemini ADK.",
 "Deploy a multi-agent system as a Streamlit application with session state.",
 "Write an MCP server with typed tools and register it with an MCP client.",
 "Explain hierarchical sub-agents and why context isolation improves reliability.",
])

h1("Assessment Preparation")
bullets([
 C.ASSESSMENT["written"],
 C.ASSESSMENT["practical"],
 "The assessment is OPEN BOOK — bring this Learner Guide and the course slides.",
 "Complete every lab; the practical tasks are based directly on what you built in class.",
 C.ASSESSMENT["note"],
 f"Submit your answers on the LMS at {C.LMS_URL}, and complete the TRAQOM survey.",
])

h1("Glossary")
gl=[
 ("Agent","An LLM that reasons in a loop, calls tools and pursues a goal, rather than answering a single prompt."),
 ("Agent skill","A modular, transferable capability — a documented procedure an agent loads when a task needs it."),
 ("Agent memory","What the agent carries forward: short-term conversation context and long-term stored facts."),
 ("Tool / function calling","A function the agent may invoke; the runtime executes it and returns the result into the loop."),
 ("MCP (Model Context Protocol)","An open standard letting any compatible agent discover and call tools exposed by a server."),
 ("MCP server","A program that publishes typed tools, resources and prompts over the Model Context Protocol."),
 ("Multi-agent system","Several specialised agents collaborating, usually coordinated by a supervisor."),
 ("Supervisor / triage agent","The agent that classifies an incoming request and routes it to the right specialist."),
 ("Handoff","Delegation of a request from one agent to another in the OpenAI Agents SDK."),
 ("Sub-agent","A child agent given a bounded task, its own tools and its own isolated context."),
 ("Context isolation","Keeping a sub-agent's working detail out of the parent's context by returning only conclusions."),
 ("Structured output","A validated, typed result (e.g. a Pydantic model) instead of free-form text."),
 ("Guardrail","An input or output check that blocks unsafe or out-of-scope work before it proceeds."),
 ("Vibe coding","Directing an AI coding agent in natural language to build software while you specify and review."),
 ("Context engineering","Deliberately curating what the coding agent sees — the main driver of output reliability."),
 ("Runner","The component that executes the agent loop (Runner.run_sync in the OpenAI SDK; Runner in ADK)."),
 ("Streamlit","A Python library for building simple web UIs, used here to deploy the agent systems."),
]
B.append(("dl",gl))

# ---------------- render Markdown ----------------
def _anchor(txt):
    return "".join(ch.lower() if ch.isalnum() else ("-" if ch in " -" else "") for ch in txt)

def render_md():
    out=[f"# {C.TITLE} — Learner Guide",""]
    out.append(f"**WSQ Course Code:** {C.COURSE_CODE}  |  **Conducted by:** {C.ORG} ({C.UEN.replace('UEN: ','UEN ')})  |  **Version {C.VERSION} · {C.VERSION_DATE}**")
    out.append("")
    out.append("## Contents"); out.append("")
    for kind,*rest in B:
        if kind=="h1": out.append(f"- [{rest[0]}](#{_anchor(rest[0])})")
        elif kind=="h2": out.append(f"  - [{rest[0]}](#{_anchor(rest[0])})")
    out.append("")
    for kind,*rest in B:
        if kind=="h1": out+=["",f"## {rest[0]}",""]
        elif kind=="h2": out+=["",f"### {rest[0]}",""]
        elif kind=="h3": out+=[f"**{rest[0]}**",""]
        elif kind=="p": out+=[rest[0],""]
        elif kind=="bullets": out+=[f"- {x}" for x in rest[0]]+[""]
        elif kind=="steps":
            for i,(instr,cmd) in enumerate(rest[0],1):
                out.append(f"{i}. {instr}")
                if cmd: out+=["","   ```bash",f"   {cmd}","   ```",""]
            out.append("")
        elif kind=="code": out+=["```bash",rest[0],"```",""]
        elif kind=="note": out+=[f"> **Note:** {rest[0]}",""]
        elif kind=="rule": out+=["---",""]
        elif kind=="dl":
            for term,defn in rest[0]: out.append(f"- **{term}** — {defn}")
            out.append("")
    return "\n".join(out)

MD_OUT=os.path.join(REPO,f"LG-{C.SHORT_TITLE}.md")
with open(MD_OUT,"w") as f: f.write(render_md())
print("Saved",MD_OUT)

# ---------------- render DOCX ----------------
BRAND=RGBColor(0x1F,0x6F,0xEB); DARK=RGBColor(0x11,0x18,0x27); GREY=RGBColor(0x55,0x5B,0x66)
INKCODE=RGBColor(0x0B,0x30,0x60)
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)
prodoc.add_cover_page(doc,"LEARNER GUIDE",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
 ("1.0",C.VERSION_DATE,
  f"Initial release — Learner Guide for AI Vibe Coding for Multi-Agents System covering all {NLABS} "
  "hands-on labs across the five topics of the published course outline.",C.TRAINER),
 ("1.1",C.VERSION_DATE,
  "Added a 'Where the code lives' section to every lab with its folder path and Google Colab link, "
  "following the restructure of the labs into one folder per lab, each holding a runnable Python "
  "script and a Colab notebook.",C.TRAINER),
])
prodoc.add_toc(doc)

def code_para(text):
    for line in text.split("\n"):
        para=doc.add_paragraph()
        r=para.add_run(line); r.font.name="Consolas"; r.font.size=Pt(9.5); r.font.color.rgb=INKCODE

for kind,*rest in B:
    if kind=="h1": doc.add_heading(rest[0],level=1)
    elif kind=="h2": doc.add_heading(rest[0],level=2)
    elif kind=="h3":
        para=doc.add_paragraph(); r=para.add_run(rest[0]); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=BRAND
    elif kind=="p": doc.add_paragraph(rest[0])
    elif kind=="bullets":
        for x in rest[0]: doc.add_paragraph(x,style="List Bullet")
    elif kind=="steps":
        # Numbers are rendered explicitly rather than via the shared "List Number" style:
        # that style continues one sequence across the whole document, so Lab 11 would start
        # at step 71 while the Markdown mirror restarts at 1. Explicit numbering keeps the
        # DOCX and the .md identical, which the house standard requires.
        for i,(instr,cmd) in enumerate(rest[0],1):
            para=doc.add_paragraph()
            para.paragraph_format.left_indent=Pt(18)
            para.paragraph_format.first_line_indent=Pt(-18)
            para.paragraph_format.space_after=Pt(4)
            r=para.add_run(f"{i}.  "); r.bold=True
            para.add_run(instr)
            if cmd: code_para(cmd)
    elif kind=="code": code_para(rest[0])
    elif kind=="note":
        para=doc.add_paragraph(); r=para.add_run("Note: "); r.bold=True; r.font.color.rgb=BRAND
        para.add_run(rest[0]).font.size=Pt(10)
    elif kind=="rule": doc.add_paragraph("")
    elif kind=="dl":
        for term,defn in rest[0]:
            para=doc.add_paragraph(style="List Bullet")
            r=para.add_run(term+" — "); r.bold=True; para.add_run(defn)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
DOCX_OUT=os.path.join(REPO,"courseware",f"LG-{C.SHORT_TITLE}.docx")
doc.save(DOCX_OUT)
print("Saved",DOCX_OUT)
