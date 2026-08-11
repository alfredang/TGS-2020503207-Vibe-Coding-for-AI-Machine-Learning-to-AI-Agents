#!/usr/bin/env python3
"""Generate the WSQ "AI Vibe Coding for Multi-Agents System" (TGS-2020503207) slide deck.

All-white Tertiary house style, driven entirely by course_data.py + data_domainN.py so the deck
stays 100% aligned with the LP, LG and labs.

DESIGN RULE FOR THIS COURSE: the deck is VISUAL — tile grids, flow diagrams, cards, architecture
diagrams and comparison panels. It deliberately contains NO per-step "step slide" walkthroughs;
the detailed step-by-step instructions live ONLY in the Learner Guide. Each lab gets a visual
overview + a workflow flow diagram + a verification slide.
"""
import os, sys, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
from data_domain1 import DOMAIN1
from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5
ACTIVITIES = DOMAIN1 + DOMAIN2 + DOMAIN3 + DOMAIN4 + DOMAIN5

def _find_repo(start):
    env = os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env):
        return env
    d = start
    for _ in range(8):
        d = os.path.dirname(d)
        if os.path.isdir(os.path.join(d, "courseware")) and os.path.isdir(os.path.join(d, "labs")):
            return d
    return os.path.dirname(os.path.dirname(HERE))
REPO = _find_repo(HERE)
ASSETS = os.path.join(os.path.dirname(HERE), "assets")

# ---------------- palette ----------------
BLUE=RGBColor(0x1F,0x6F,0xEB); TEAL=RGBColor(0x10,0xB9,0x81); AMBER=RGBColor(0xF5,0x9E,0x0B)
INK=RGBColor(0x16,0x1B,0x26); GREY=RGBColor(0x5B,0x63,0x72); LIGHT=RGBColor(0xF5,0xF8,0xFC)
WHITE=RGBColor(0xFF,0xFF,0xFF); LINE=RGBColor(0xE2,0xE8,0xF0); VIOLET=RGBColor(0x7C,0x3A,0xED)
ROSE=RGBColor(0xE1,0x1D,0x48)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height
BLANK=prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)
def rect(s,x,y,w,h,color,line=None):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(1)
    sp.shadow.inherit=False; return sp
def rrect(s,x,y,w,h,color,line=None):
    sp=s.shapes.add_shape(5,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(1.25)
    sp.shadow.inherit=False; return sp
def oval(s,x,y,w,h,color):
    sp=s.shapes.add_shape(9,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    sp.line.fill.background(); sp.shadow.inherit=False; return sp
def arrow(s,x,y,w,h,color):
    sp=s.shapes.add_shape(13,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,space=4):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,line in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(space)
        for t,sz,col,bold in line:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=bold
            r.font.color.rgb=col; r.font.name="Arial"
    return tb
def bullets(s,x,y,w,h,items,size=18,color=INK,gap=10,mcolor=BLUE):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(gap)
        lvl=it[1] if isinstance(it,tuple) else 0
        text=it[0] if isinstance(it,tuple) else it
        r=p.add_run(); r.text=("•  " if lvl==0 else "–  ")+text
        r.font.size=Pt(size if lvl==0 else size-2); r.font.color.rgb=color if lvl==0 else GREY
        r.font.name="Arial"; r.font.bold=(lvl==0 and isinstance(it,tuple) and len(it)>2 and it[2])
    return tb

PAGE={"n":0}
def footer(s):
    PAGE["n"]+=1
    txt(s,Inches(0.4),Inches(7.05),Inches(7.5),Inches(0.35),
        [[(f"{C.SHORT_TITLE}  ·  {C.COURSE_CODE}",9,GREY,False)]])
    txt(s,Inches(5.0),Inches(7.05),Inches(3.3),Inches(0.35),
        [[("© 2026 Tertiary Infotech Academy Pte Ltd",9,GREY,False)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(12.4),Inches(7.05),Inches(0.6),Inches(0.35),
        [[(str(PAGE["n"]),9,GREY,False)]],align=PP_ALIGN.RIGHT)
def head(s,title,kicker=None,kcolor=BLUE):
    rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),Inches(1.55),kcolor)
    if kicker: txt(s,Inches(0.85),Inches(0.5),Inches(11.6),Inches(0.4),[[(kicker,14,kcolor,True)]])
    txt(s,Inches(0.85),Inches(0.9),Inches(11.9),Inches(0.9),[[(title,29,INK,True)]])
    rect(s,Inches(0.85),Inches(1.7),Inches(11.63),Inches(0.02),LINE)
    return s
def _logo(name):
    p=os.path.join(ASSETS,name)
    return p if os.path.exists(p) else None

# ---------------- slide templates ----------------
def cover():
    s=slide(); rect(s,0,0,SW,SH,WHITE)
    rect(s,0,0,SW,Inches(0.22),BLUE); rect(s,0,Inches(7.28),SW,Inches(0.22),TEAL)
    org=_logo("tertiary-infotech-logo.png")
    if org: s.shapes.add_picture(org,Inches(0.85),Inches(0.7),height=Inches(1.05))
    # course badge (top-right) — multi-agent motif
    bx=Inches(10.55); by=Inches(0.55); bw=Inches(2.0); bh=Inches(1.45)
    rrect(s,bx,by,bw,bh,LIGHT,line=BLUE)
    # three connected nodes = multi-agent motif
    cx=bx+Inches(0.35); cy=by+Inches(0.26)
    rect(s,cx+Inches(0.2),cy+Inches(0.16),Inches(0.5),Inches(0.03),LINE)
    rect(s,cx+Inches(0.8),cy+Inches(0.16),Inches(0.5),Inches(0.03),LINE)
    oval(s,cx+Inches(0.5),cy,Inches(0.32),Inches(0.32),BLUE)
    oval(s,cx-Inches(0.02),cy+Inches(0.44),Inches(0.32),Inches(0.32),TEAL)
    oval(s,cx+Inches(1.02),cy+Inches(0.44),Inches(0.32),Inches(0.32),VIOLET)
    txt(s,bx,by+Inches(1.02),bw,Inches(0.32),[[("MULTI-AGENT",9,GREY,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(0.9),Inches(2.3),Inches(12),Inches(0.6),[[("COURSE SLIDES  ·  WSQ",16,BLUE,True)]])
    txt(s,Inches(0.9),Inches(2.85),Inches(11.4),Inches(1.9),[[(C.TITLE,40,INK,True)]])
    rect(s,Inches(0.92),Inches(4.75),Inches(2.4),Inches(0.06),TEAL)
    txt(s,Inches(0.9),Inches(5.05),Inches(12),Inches(1.4),
        [[(f"WSQ Course Code: {C.COURSE_CODE}",16,GREY,False)],
         [(f"Conducted by {C.ORG}  ·  {C.UEN.replace('UEN: ','UEN ')}",14,GREY,False)],
         [(f"Trainer: {C.TRAINER}",14,GREY,False)]],space=6)
    txt(s,Inches(0.9),Inches(6.5),Inches(12),Inches(0.4),[[(f"Version {C.VERSION}  ·  {C.VERSION_DATE}",12,GREY,False)]])
    txt(s,Inches(0.9),Inches(6.85),Inches(12),Inches(0.34),[[("© 2026 Tertiary Infotech Academy Pte Ltd. All rights reserved.  ·  www.tertiarycourses.com.sg",10,GREY,False)]])

def section(kicker,title,n,sub=""):
    s=slide(); rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),SH,BLUE)
    rect(s,Inches(0.85),Inches(2.5),Inches(0.14),Inches(2.0),TEAL)
    txt(s,Inches(1.25),Inches(2.55),Inches(11),Inches(0.6),[[(kicker,18,BLUE,True)]])
    txt(s,Inches(1.25),Inches(3.0),Inches(10.6),Inches(1.6),[[(title,40,INK,True)]])
    if sub: txt(s,Inches(1.27),Inches(4.55),Inches(11),Inches(0.8),[[(sub,16,GREY,False)]])
    txt(s,Inches(10.0),Inches(0.7),Inches(2.8),Inches(1.6),[[(n,72,RGBColor(0xE2,0xE8,0xF0),True)]],align=PP_ALIGN.RIGHT)
    footer(s)
def content(title,items,kicker=None,size=20):
    s=head(slide(),title,kicker); bullets(s,Inches(0.85),Inches(1.95),Inches(11.6),Inches(4.9),items,size=size); footer(s); return s
def two_col(title,left,right,kicker=None,lhead="",rhead=""):
    s=head(slide(),title,kicker)
    rect(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(4.7),LIGHT); rect(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(4.7),LIGHT)
    rect(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(0.1),BLUE); rect(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(0.1),TEAL)
    if lhead: txt(s,Inches(1.1),Inches(2.2),Inches(5.2),Inches(0.4),[[(lhead,16,BLUE,True)]])
    if rhead: txt(s,Inches(7.2),Inches(2.2),Inches(5.0),Inches(0.4),[[(rhead,16,TEAL,True)]])
    bullets(s,Inches(1.1),Inches(2.75),Inches(5.2),Inches(3.7),left,size=15)
    bullets(s,Inches(7.2),Inches(2.75),Inches(5.05),Inches(3.7),right,size=15,mcolor=TEAL); footer(s); return s
def cards3(title,cards,kicker):
    s=head(slide(),title,kicker); xs=[Inches(0.85),Inches(5.0),Inches(9.15)]
    for i,c in enumerate(cards[:3]):
        x=xs[i]; col=c[0]
        rect(s,x,Inches(1.95),Inches(3.65),Inches(4.7),LIGHT); rect(s,x,Inches(1.95),Inches(3.65),Inches(0.12),col)
        txt(s,x+Inches(0.25),Inches(2.2),Inches(3.2),Inches(0.6),[[(c[1],19,col,True)]])
        bullets(s,x+Inches(0.25),Inches(2.95),Inches(3.2),Inches(3.4),c[2],size=14,mcolor=col,gap=9)
    footer(s); return s
def big_statement(line1,line2,kicker,color=BLUE):
    s=slide(); rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),SH,color)
    txt(s,Inches(1.1),Inches(2.2),Inches(11),Inches(0.5),[[(kicker,16,color,True)]])
    txt(s,Inches(1.1),Inches(2.8),Inches(11.3),Inches(2.4),[[(line1,38,INK,True)]])
    if line2: txt(s,Inches(1.12),Inches(4.9),Inches(11),Inches(1.2),[[(line2,20,GREY,False)]])
    footer(s); return s

PALETTE=[BLUE,TEAL,VIOLET,AMBER]
def tile_grid(title,items,kicker=None,cols=2,size=15,icons=None,accent=BLUE):
    """Grid of light panels, each with a coloured icon/number badge + text."""
    s=head(slide(),title,kicker,kcolor=accent)
    n=len(items); rows=math.ceil(n/cols)
    X0=Inches(0.85); Y0=Inches(1.95); TOTW=Inches(11.63); AREAH=Inches(4.78)
    gx=Inches(0.3); gy=Inches(0.26)
    cw=int((TOTW-gx*(cols-1))/cols); ch=int((AREAH-gy*(rows-1))/rows)
    bd=Inches(0.6)
    for i,it in enumerate(items):
        r=i//cols; c=i%cols
        x=int(X0+(cw+gx)*c); y=int(Y0+(ch+gy)*r); col=PALETTE[i%len(PALETTE)]
        rect(s,x,y,cw,ch,LIGHT); rect(s,x,y,Inches(0.1),ch,col)
        oval(s,x+Inches(0.28),int(y+ch/2-bd/2),bd,bd,col)
        ic=icons[i] if icons else str(i+1)
        txt(s,x+Inches(0.28),int(y+ch/2-bd/2),bd,bd,[[(ic,19,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        tx=x+Inches(1.08); tw=cw-Inches(1.32)
        if isinstance(it,tuple):
            txt(s,tx,int(y+Inches(0.14)),tw,int(ch-Inches(0.2)),
                [[(it[0],size+2,INK,True)],[(it[1],size-2,GREY,False)]],anchor=MSO_ANCHOR.MIDDLE,space=3)
        else:
            txt(s,tx,int(y+Inches(0.1)),tw,int(ch-Inches(0.16)),[[(it,size,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def flow_h(title,steps,kicker=None,color=BLUE):
    """Horizontal numbered flow: coloured chips connected by chevrons."""
    s=head(slide(),title,kicker,kcolor=color)
    n=len(steps); X0=Inches(0.85); TOTW=Inches(11.63); gap=Inches(0.34)
    cw=int((TOTW-gap*(n-1))/n); y=Inches(2.45); ch=Inches(2.85); bd=Inches(0.78)
    for i,st in enumerate(steps):
        x=int(X0+(cw+gap)*i)
        rect(s,x,y,cw,ch,LIGHT); rect(s,x,y,cw,Inches(0.1),color)
        oval(s,int(x+cw/2-bd/2),int(y+Inches(0.38)),bd,bd,color)
        txt(s,int(x+cw/2-bd/2),int(y+Inches(0.38)),bd,bd,[[(str(i+1),28,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        txt(s,x+Inches(0.14),int(y+Inches(1.32)),cw-Inches(0.28),int(ch-Inches(1.45)),[[(st,13,INK,False)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.TOP)
        if i<n-1:
            txt(s,int(x+cw-Inches(0.04)),int(y+ch/2-Inches(0.3)),int(gap+Inches(0.08)),Inches(0.6),
                [[("▶",15,color,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s

def agent_loop_diagram(title,kicker):
    """The reason-act-observe loop drawn as a cycle of four nodes."""
    s=head(slide(),title,kicker,kcolor=BLUE)
    nodes=[("GOAL","The user's request enters the loop",BLUE),
           ("REASON","The model decides the next action",VIOLET),
           ("ACT","A tool is called and executed",TEAL),
           ("OBSERVE","The result re-enters the context",AMBER)]
    X0=Inches(0.95); y=Inches(2.7); bw=Inches(2.5); bh=Inches(2.1); gap=Inches(0.52)
    for i,(name,desc,col) in enumerate(nodes):
        x=int(X0+(bw+gap)*i)
        rrect(s,x,y,bw,bh,LIGHT,line=col)
        rect(s,x,y,bw,Inches(0.11),col)
        txt(s,x,y+Inches(0.42),bw,Inches(0.5),[[(name,18,col,True)]],align=PP_ALIGN.CENTER)
        txt(s,x+Inches(0.16),y+Inches(1.0),bw-Inches(0.32),Inches(0.95),[[(desc,12,GREY,False)]],align=PP_ALIGN.CENTER)
        if i<len(nodes)-1:
            arrow(s,int(x+bw+Inches(0.08)),int(y+bh/2-Inches(0.13)),int(gap-Inches(0.16)),Inches(0.26),col)
    # feedback arrow back to REASON
    rect(s,int(X0+Inches(1.2)),Inches(5.15),int(bw*3+gap*3-Inches(0.5)),Inches(0.035),GREY)
    txt(s,Inches(0.95),Inches(5.25),Inches(11.6),Inches(0.4),
        [[("↖  the loop repeats until the goal is met or the turn limit is reached",12,GREY,True)]],align=PP_ALIGN.CENTER)
    footer(s); return s

def architecture_diagram(title,kicker,supervisor,specialists,tools_label="TOOLS / MCP SERVERS",accent=BLUE):
    """Supervisor → specialists → tools, drawn as a three-tier hierarchy."""
    s=head(slide(),title,kicker,kcolor=accent)
    # tier 1 — supervisor
    sx=Inches(4.9); sw=Inches(3.5); sy=Inches(2.05); sh=Inches(0.85)
    rrect(s,sx,sy,sw,sh,accent)
    txt(s,sx,sy,sw,sh,[[(supervisor,17,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    # tier 2 — specialists
    n=len(specialists); X0=Inches(0.9); TOTW=Inches(11.55); gap=Inches(0.3)
    cw=int((TOTW-gap*(n-1))/n); y2=Inches(3.55); h2=Inches(1.5)
    cols=[TEAL,VIOLET,AMBER,ROSE]
    for i,(nm,desc) in enumerate(specialists):
        x=int(X0+(cw+gap)*i); col=cols[i%len(cols)]
        # connector from supervisor
        rect(s,int(x+cw/2),Inches(2.98),Inches(0.03),Inches(0.55),LINE)
        rrect(s,x,y2,cw,h2,LIGHT,line=col)
        rect(s,x,y2,cw,Inches(0.1),col)
        txt(s,x+Inches(0.12),y2+Inches(0.28),cw-Inches(0.24),Inches(0.45),[[(nm,14,col,True)]],align=PP_ALIGN.CENTER)
        txt(s,x+Inches(0.12),y2+Inches(0.75),cw-Inches(0.24),Inches(0.65),[[(desc,11,GREY,False)]],align=PP_ALIGN.CENTER)
        rect(s,int(x+cw/2),Inches(5.05),Inches(0.03),Inches(0.45),LINE)
    # tier 3 — tools
    ty=Inches(5.5); th=Inches(0.8)
    rect(s,Inches(0.9),ty,Inches(11.55),th,LIGHT); rect(s,Inches(0.9),ty,Inches(0.1),th,GREY)
    txt(s,Inches(0.9),ty,Inches(11.55),th,[[(tools_label,14,GREY,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s

def compare_panel(title,kicker,left_title,left_items,right_title,right_items):
    """Two-column comparison with coloured heads — for single vs multi-agent, SDK vs SDK."""
    s=head(slide(),title,kicker)
    rrect(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(4.7),LIGHT,line=BLUE)
    rrect(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(4.7),LIGHT,line=TEAL)
    rect(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(0.5),BLUE)
    rect(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(0.5),TEAL)
    txt(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(0.5),[[(left_title,16,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(0.5),[[(right_title,16,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    bullets(s,Inches(1.15),Inches(2.7),Inches(5.15),Inches(3.7),left_items,size=14)
    bullets(s,Inches(7.25),Inches(2.7),Inches(4.95),Inches(3.7),right_items,size=14,mcolor=TEAL)
    footer(s); return s

def tool_landscape(title,kicker,terminal,ide):
    """Vibe coding tool landscape — two labelled groups of tool chips."""
    s=head(slide(),title,kicker,kcolor=VIOLET)
    for gi,(label,group,col) in enumerate([("TERMINAL AGENTS",terminal,BLUE),("AGENTIC IDEs",ide,TEAL)]):
        y=Inches(2.05)+gi*Inches(2.45)
        rect(s,Inches(0.85),y,Inches(0.1),Inches(2.05),col)
        txt(s,Inches(1.1),y,Inches(3.0),Inches(0.4),[[(label,13,col,True)]])
        n=len(group); X0=Inches(1.1); cw=Inches(2.6); gap=Inches(0.32)
        for i,(name,url) in enumerate(group):
            x=int(X0+(cw+gap)*i)
            rrect(s,x,int(y+Inches(0.5)),cw,Inches(1.35),LIGHT,line=col)
            txt(s,x+Inches(0.12),int(y+Inches(0.68)),cw-Inches(0.24),Inches(0.45),[[(name,15,INK,True)]],align=PP_ALIGN.CENTER)
            txt(s,x+Inches(0.1),int(y+Inches(1.18)),cw-Inches(0.2),Inches(0.5),[[(url.replace("https://",""),9,GREY,False)]],align=PP_ALIGN.CENTER)
    footer(s); return s

def trainer_slide(kicker,name,role,rows,initials,accent=BLUE):
    s=head(slide(),"About the Trainer",kicker,kcolor=accent)
    lx=Inches(0.85); lw=Inches(3.65)
    rect(s,lx,Inches(1.95),lw,Inches(4.7),LIGHT); rect(s,lx,Inches(1.95),lw,Inches(0.12),accent)
    bd=Inches(1.7); ax=int(lx+(lw-bd)/2)
    oval(s,ax,Inches(2.5),bd,bd,accent)
    txt(s,ax,Inches(2.5),bd,bd,[[(initials,44,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(s,lx+Inches(0.15),Inches(4.55),lw-Inches(0.3),Inches(0.6),[[(name,21,INK,True)]],align=PP_ALIGN.CENTER)
    txt(s,lx+Inches(0.15),Inches(5.2),lw-Inches(0.3),Inches(1.2),[[(role,13,GREY,False)]],align=PP_ALIGN.CENTER)
    rx=Inches(4.9); rw=Inches(7.6); ry=Inches(1.95); rh=Inches(4.7)
    n=len(rows); gy=Inches(0.2); th=int((rh-gy*(n-1))/n)
    for i,(label,val) in enumerate(rows):
        y=int(ry+(th+gy)*i); col=PALETTE[i%len(PALETTE)]
        rect(s,rx,y,rw,th,LIGHT); rect(s,rx,y,Inches(0.1),th,col)
        vruns=[(val,14,INK,False)] if val else [("____________________________________________",13,LINE,False)]
        txt(s,rx+Inches(0.32),y,rw-Inches(0.6),th,
            [[(label.upper(),11,col,True)],vruns],anchor=MSO_ANCHOR.MIDDLE,space=3)
    footer(s); return s

def activity_overview(tag,title,desc,build,services,kicker,lab_path=None):
    s=head(slide(),title,kicker,kcolor=TEAL)
    rect(s,Inches(0.85),Inches(1.85),Inches(1.7),Inches(0.5),TEAL)
    txt(s,Inches(0.85),Inches(1.9),Inches(1.7),Inches(0.4),[[(tag,16,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(0.85),Inches(2.5),Inches(11.7),Inches(1.5),[[(desc,20,INK,False)]])
    rect(s,Inches(0.85),Inches(4.1),Inches(11.7),Inches(1.75),LIGHT)
    rect(s,Inches(0.85),Inches(4.1),Inches(0.1),Inches(1.75),BLUE)
    txt(s,Inches(1.1),Inches(4.28),Inches(11),Inches(0.4),[[("You'll build",14,BLUE,True)]])
    txt(s,Inches(1.1),Inches(4.66),Inches(11.2),Inches(0.6),[[(build,17,INK,True)]])
    txt(s,Inches(1.1),Inches(5.36),Inches(11.2),Inches(0.5),[[("Tools:  ",13,GREY,True),(services,13,GREY,False)]])
    # Where the code lives — the lab folder and its Colab notebook.
    if lab_path:
        rrect(s,Inches(0.85),Inches(6.0),Inches(11.7),Inches(0.62),WHITE,line=TEAL)
        rect(s,Inches(0.85),Inches(6.0),Inches(0.1),Inches(0.62),TEAL)
        txt(s,Inches(1.15),Inches(6.0),Inches(11.2),Inches(0.62),
            [[("Code:  ",12,TEAL,True),(f"labs/{lab_path}/",12,INK,False),
              ("     ·     Open in Colab from the lab README or the course repo",12,GREY,False)]],
            anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s

def lab_workflow(act):
    """VISUAL summary of the lab's workflow — NOT a step-by-step walkthrough.
    The detailed steps live in the Learner Guide only.

    Each lab supplies its own short phase labels in `phases` (4-5 items, <= ~40 chars each)
    so the slide stays a readable diagram. Labels are authored, never truncated from the
    step text — a clipped mid-word label is a house-standard defect."""
    phases=act["phases"]
    flow_h(f"Lab {act['num']} Workflow",phases,
           kicker=f"LAB {act['num']} · AT A GLANCE  ·  FULL STEPS IN THE LEARNER GUIDE",color=TEAL)

def test_slide(act_title,text,kicker):
    s=head(slide(),act_title,kicker,TEAL)
    rect(s,Inches(0.85),Inches(2.3),Inches(11.7),Inches(2.6),RGBColor(0xE8,0xF7,0xEE))
    rect(s,Inches(0.85),Inches(2.3),Inches(0.1),Inches(2.6),TEAL)
    txt(s,Inches(1.2),Inches(2.6),Inches(11),Inches(0.5),[[("✅  Test it",20,RGBColor(0x12,0x7A,0x3E),True)]])
    txt(s,Inches(1.2),Inches(3.3),Inches(11),Inches(1.4),[[(text,17,INK,False)]])
    txt(s,Inches(0.85),Inches(5.2),Inches(11.7),Inches(0.5),
        [[("Detailed step-by-step instructions for this lab are in the Learner Guide.",13,GREY,False)]])
    footer(s); return s

def brk(kind,dur,color=AMBER):
    s=slide(); rect(s,0,0,SW,SH,WHITE)
    rect(s,0,0,SW,Inches(0.22),color); rect(s,0,Inches(7.28),SW,Inches(0.22),color)
    rect(s,Inches(5.4),Inches(2.35),Inches(2.53),Inches(0.1),color)
    txt(s,0,Inches(2.75),SW,Inches(1.2),[[(kind,48,INK,True)]],align=PP_ALIGN.CENTER)
    txt(s,0,Inches(4.05),SW,Inches(0.8),[[(dur,22,color,True)]],align=PP_ALIGN.CENTER); PAGE["n"]+=1

def download_material_slide():
    """Visual 'Download Course Material' slide — numbered portal steps as a flow."""
    s=head(slide(),"Download Your Course Material","COURSE PORTAL · LMS/TMS",kcolor=BLUE)
    # browser chrome mock
    bx=Inches(0.85); by=Inches(2.0); bw=Inches(6.4); bh=Inches(4.35)
    rrect(s,bx,by,bw,bh,WHITE,line=LINE)
    rect(s,bx,by,bw,Inches(0.46),LIGHT)
    for i,c in enumerate([ROSE,AMBER,TEAL]):
        oval(s,bx+Inches(0.2)+i*Inches(0.26),by+Inches(0.15),Inches(0.16),Inches(0.16),c)
    rrect(s,bx+Inches(1.05),by+Inches(0.1),bw-Inches(1.35),Inches(0.27),WHITE,line=LINE)
    txt(s,bx+Inches(1.2),by+Inches(0.09),bw-Inches(1.5),Inches(0.3),
        [[("lms-tms.tertiaryinfotech.com",10,GREY,False)]],anchor=MSO_ANCHOR.MIDDLE)
    txt(s,bx+Inches(0.3),by+Inches(0.75),bw-Inches(0.6),Inches(0.4),[[("My Courses",15,INK,True)]])
    for i,(lbl,col) in enumerate([("Trainer Slides (PPT)",BLUE),("Learner Slides (PDF)",TEAL),
                                   ("Learner Guide (PDF)",VIOLET),("Lesson Plan (PDF)",AMBER)]):
        ry=by+Inches(1.3)+i*Inches(0.72)
        rect(s,bx+Inches(0.3),ry,bw-Inches(0.6),Inches(0.58),LIGHT)
        rect(s,bx+Inches(0.3),ry,Inches(0.08),Inches(0.58),col)
        txt(s,bx+Inches(0.55),ry,Inches(3.6),Inches(0.58),[[(lbl,12,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
        rrect(s,bx+bw-Inches(1.35),ry+Inches(0.11),Inches(1.0),Inches(0.36),col)
        txt(s,bx+bw-Inches(1.35),ry+Inches(0.11),Inches(1.0),Inches(0.36),
            [[("Download",9,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    # steps beside it
    for i,st in enumerate(["Go to lms-tms.tertiaryinfotech.com and sign in with the email you registered with.",
                            "Open My Courses and select this course.",
                            "Download the Trainer Slides, Learner Guide and Lesson Plan.",
                            "Keep them open — the assessment is OPEN BOOK."]):
        y=Inches(2.15)+i*Inches(1.06)
        col=PALETTE[i%len(PALETTE)]
        oval(s,Inches(7.6),y,Inches(0.5),Inches(0.5),col)
        txt(s,Inches(7.6),y,Inches(0.5),Inches(0.5),[[(str(i+1),16,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        txt(s,Inches(8.3),y-Inches(0.02),Inches(4.2),Inches(0.95),[[(st,13,INK,False)]])
    footer(s); return s

# ============================================================ BUILD
cover()

# ---------------- ADMIN ----------------
section("COURSE ADMINISTRATION","Welcome & Housekeeping","")
flow_h("Digital Attendance (Mandatory)",[
 "The trainer displays the SSG digital attendance QR code",
 "Scan the QR code with your mobile phone camera",
 "Submit your attendance — AM, PM and Assessment",
 "Repeat on every training day",
 "75% attendance is required for assessment and funding"],
 kicker="TRAQOM · SSG DIGITAL ATTENDANCE",color=AMBER)
trainer_slide("YOUR TRAINER · GENERAL","Your Trainer","General Trainer template —\nto be completed by the trainer",
 [("Name",""),("Title / Designation",""),("Qualifications",""),
  ("Areas of expertise",""),("Training & industry experience",""),("Contact","")],
 initials="?",accent=GREY)
trainer_slide("YOUR TRAINER",C.TRAINER,"Principal Trainer\nTertiary Infotech Academy Pte Ltd",
 [("Role","Principal Trainer, Tertiary Infotech Academy Pte Ltd"),
  ("Expertise","AI agents, multi-agent systems, machine learning and full-stack AI application development."),
  ("Delivers","WSQ courses on AI vibe coding, agentic AI, machine learning and data analytics."),
  ("Profile","Founder and lead instructor at Tertiary Infotech / Tertiary Courses  ·  github.com/alfredang")],
 initials="AA",accent=BLUE)
content("Let's Know Each Other",[
 "Your name, organisation and role.",
 "Your experience with Python, LLMs or AI coding tools (if any).",
 "What you want to be able to build with AI agents after this course."],kicker="ICE-BREAKER")
tile_grid("Ground Rules",[
 "Set your mobile phone to silent mode.","Participate actively — no question is too small.",
 "Mutual respect: agree to disagree.","One conversation at a time.",
 "Be punctual; return from breaks on time.","75% attendance is required for funding."],
 kicker="HOUSEKEEPING",cols=2,size=15)
download_material_slide()
tile_grid("Skills Framework Alignment",
 [(C.TSC_TITLE,f"TSC Code: {C.TSC_CODE}")]+[(f"Ability {a.split(' - ')[0]}",a.split(' - ',1)[1]) for a in C.TSC_ABILITIES],
 kicker="TSC MAPPING",cols=1,size=13)
# Lesson plan overview (2 days)
two_col("Lesson Plan — 2 Days, 8 Hours per Day",[
 (f"Day 1 — {C.DAY_THEMES[1]}",0),
 ("Digital Attendance (AM), introductions, learning outcomes",1),
 ("Topic 1: Modern Agent Foundations (Labs 1–2)",1),
 ("Topic 2: Vibe Coding for Multi-Agent Systems (Labs 3–4)",1),
 ("Topic 3 begins: OpenAI Agents SDK (Lab 5)",1),
 ("Day 1 recap and PM digital attendance",1)],
 [(f"Day 2 — {C.DAY_THEMES[2]}",0),
 ("Topic 3 continues: supervisor routing, Streamlit (Labs 6–7)",1),
 ("Topic 4: Gemini Agent SDK (Labs 8–9)",1),
 ("Topic 5: MCP and Sub-Agents (Labs 10–11)",1),
 ("Course feedback and TRAQOM survey",1),
 ("Final Assessment: WA (SAQ) + PP",1),
 ("Daily timing: 9:30am–7:00pm · 8 instructional hours · 1-hour lunch · 2 tea breaks",1)],
 kicker="SCHEDULE",lhead="Day 1",rhead="Day 2 & timing")
tile_grid("Learning Outcomes",[
 ("LO1 · Agent Foundations","Skills, memory, tools, MCP, and single-agent to multi-agent progression."),
 ("LO2 · Vibe Coding","A structured workflow with context engineering for reliable agent code."),
 ("LO3 · OpenAI Agents SDK","Structured outputs, tool calling, supervisor routing, Streamlit deployment."),
 ("LO4 · Gemini Agent SDK","Collaborative Gemini agents, deployed with Streamlit."),
 ("LO5 · MCP & Sub-Agents","Tool orchestration and hierarchical sub-agent architecture.")],
 kicker="WHAT YOU'LL ACHIEVE",cols=1,size=14)
tile_grid("Course Outline",[(f"Topic {t['code']} — {t['title']}",t["subtitle"]) for t in C.TOPICS],
 kicker="FIVE TOPICS · ELEVEN HANDS-ON LABS",cols=1,size=14)
tile_grid("Briefing for Assessment",[
 ("Phones away","Place phones and other materials under the table or on the floor."),
 ("No recording","No photos or recording of assessment scripts."),
 ("Work alone","No discussion of any kind during the assessment."),
 ("Black or blue pen","Use a black/blue pen for hard-copy assessments."),
 ("No correction fluid","No liquid paper or correction tape on the script."),
 ("Time is called","Scripts are collected when time is up.")],
 kicker="BEFORE THE ASSESSMENT",cols=2,size=14,accent=AMBER)
cards3("Assessment",[
 (BLUE,"Written (WA)",["Short-Answer Questions (SAQ)","50 minutes","Tests underpinning knowledge",
                       "Answer in your own words"]),
 (TEAL,"Practical (PP)",["Hands-on multi-agent build tasks","75 minutes",
                         "Tests practical ability","Based on the labs you did in class"]),
 (VIOLET,"Rules",["Open book — slides, Learner Guide and approved materials",
                   "75% attendance required","Graded Competent / Not Yet Competent",
                   "An appeal process is available"])],kicker="FINAL ASSESSMENT")
flow_h("Assessment Flow",[
 "TRAQOM survey — scan the QR code on the LMS",
 "Assessment digital attendance — scan the SSG QR",
 "Sit WA (SAQ) then PP — open book",
 "Submit your answers on the LMS",
 "Sign the Assessment Summary Record"],kicker="ON ASSESSMENT DAY")
tile_grid("Codes for the Labs",[
 ("GitHub repository","github.com/tertiarycourses/TGS-2020503207-AI-Vibe-Coding-for-Multi-Agents-System"),
 ("One folder per lab","labs/lab-01-… to labs/lab-11-… — each with a runnable Python script and a README."),
 ("Run it locally","python <script>.py with your virtual environment activated and your .env keys set."),
 ("Or open in Colab","Every lab ships a notebook — click the Open in Colab badge in its README."),
 ("Learner Guide","The full detailed step-by-step for every lab — your open-book reference.")],
 kicker="LAB RESOURCES",cols=1,size=14)

# ---------------- CORE CONCEPTS ----------------
section("CORE CONCEPTS","From One Agent to Many","")
tile_grid("What is an AI Agent?",[
 ("Not just a chatbot","A chatbot answers; an agent pursues a goal, takes actions and checks its own results."),
 ("The reasoning loop","Reason, act, observe — repeated until the goal is met or a limit is reached."),
 ("Tools","Functions the agent may call to reach beyond the model: APIs, files, databases, the shell."),
 ("Memory","What it carries forward — the conversation now, and durable facts across sessions."),
 ("Autonomy with limits","The agent chooses the path; you set the boundaries, the budget and the review gate."),
 ("Why now","Reliable tool calling and long context finally make the loop dependable enough to ship.")],
 kicker="FOUNDATIONS",cols=2,size=14)
agent_loop_diagram("The Agent Reasoning Loop","HOW AN AGENT ACTUALLY WORKS")
compare_panel("Single Agent vs Multi-Agent System","WHEN TO SPLIT",
 "SINGLE AGENT",[("Best when",0),("The task is narrow and well defined",1),("A handful of tools is enough",1),
  ("Context stays small",1),("Limits",0),("Too many tools confuse tool choice",1),
  ("One long context mixes unrelated work",1),("One instruction set must cover everything",1)],
 "MULTI-AGENT SYSTEM",[("Best when",0),("The work splits into distinct specialities",1),
  ("Each part needs its own tools and rules",1),("Parts can run in parallel",1),
  ("Benefits",0),("Focused instructions per agent",1),("Isolated context per agent",1),
  ("Easier to test and replace one part",1)])
architecture_diagram("Multi-Agent Architecture Pattern","THE PATTERN YOU WILL BUILD TODAY","SUPERVISOR / TRIAGE AGENT",
 [("Research Agent","Searches and gathers sources"),("Coding Agent","Writes and runs code"),
  ("Writing Agent","Drafts and edits prose"),("Analysis Agent","Computes and interprets")])
big_statement("One agent that does everything does nothing well.","Split the work by speciality, give each agent its own tools and context, and let a supervisor route the request.","WHY MULTI-AGENT",color=BLUE)
tile_grid("Model Context Protocol (MCP)",[
 ("The problem","Every framework needed its own bespoke integration for the same tool."),
 ("The standard","MCP is an open protocol for exposing tools, resources and prompts to any client."),
 ("Servers","A server publishes typed tools — you write the capability once."),
 ("Clients","Any MCP-aware agent discovers those tools and calls them."),
 ("Transport","stdio for local servers; HTTP/SSE for remote ones."),
 ("The payoff","Write it once, and every MCP-aware agent can use it.")],
 kicker="THE UNIVERSAL TOOL INTERFACE",cols=2,size=14,accent=VIOLET)

# ---------------- TOPICS + ACTIVITIES ----------------
TOPIC_ACTS = {t["num"]: [a for a in ACTIVITIES if a["topic"]==t["num"]] for t in C.TOPICS}
CARD_COLORS=[BLUE,TEAL,VIOLET]
for t in C.TOPICS:
    section(f"TOPIC {t['code']}", t["title"], t["code"], t["subtitle"])
    tile_grid(f"Key Concepts — {t['title']}", t["concepts"],
              kicker=f"TOPIC {t['code']}  ·  {t['weighting']} OF THE COURSE", cols=2, size=13)

    # topic-specific visual explainers
    if t["num"]==2:
        flow_h("The Structured Vibe Coding Workflow",
               ["SPECIFY — write the spec and the acceptance test first",
                "SCAFFOLD — let the agent build the structure",
                "GENERATE — implement one component at a time",
                "RUN & VERIFY — execute, paste real errors back",
                "REFINE — one improvement per iteration, commit each"],
               kicker="SPECIFY → SCAFFOLD → GENERATE → VERIFY → REFINE",color=VIOLET)
        tool_landscape("Vibe Coding Tools","CHOOSE YOUR AGENT",C.VIBE_TOOLS_TERMINAL,C.VIBE_TOOLS_IDE)
        tile_grid("Context Engineering — What Actually Drives Reliability",[
         ("Curate, don't dump","Show the agent the few files that matter, not the whole repository."),
         ("Project context file","CLAUDE.md / GEMINI.md states the stack, conventions and prohibitions once."),
         ("Precise references","Point at file paths and line ranges instead of pasting large blocks."),
         ("Skills","Package a repeatable procedure so quality does not depend on re-explaining it."),
         ("Specify before generating","A written spec and acceptance test beat any amount of prompt wording."),
         ("Stay the reviewer","Read every diff. You own the architecture and the acceptance decision.")],
         kicker="THE HIGHEST-LEVERAGE SKILL",cols=2,size=13,accent=VIOLET)
    if t["num"]==3:
        architecture_diagram("OpenAI Agents SDK — Triage and Handoffs","TOPIC 03 · ARCHITECTURE","TRIAGE AGENT  (handoffs)",
         [("Research Agent","Web search tools"),("Coding Agent","Code execution tools"),("Writing Agent","Drafting and editing")],
         tools_label="@function_tool  ·  PYDANTIC STRUCTURED OUTPUTS  ·  GUARDRAILS",accent=BLUE)
        tile_grid("The Building Blocks",[
         ("Agent","A model plus instructions plus tools — the unit of specialisation."),
         ("Runner","Executes the agent loop: Runner.run_sync(agent, prompt)."),
         ("@function_tool","Turns a Python function into a callable tool from its type hints and docstring."),
         ("output_type","A Pydantic model that forces validated, typed output instead of free prose."),
         ("handoffs","The list of specialists a triage agent may delegate to."),
         ("Guardrails","Input and output checks that stop out-of-scope or unsafe work early.")],
         kicker="OPENAI AGENTS SDK",cols=2,size=13)
    if t["num"]==4:
        architecture_diagram("Gemini Agent SDK (ADK) — Coordinator and Sub-Agents","TOPIC 04 · ARCHITECTURE","COORDINATOR AGENT  (sub_agents)",
         [("Specialist A","Own tools and instruction"),("Specialist B","Own tools and instruction"),("Specialist C","Own tools and instruction")],
         tools_label="PYTHON FUNCTION TOOLS  ·  RUNNER  ·  SESSION SERVICE",accent=TEAL)
        compare_panel("OpenAI Agents SDK vs Gemini ADK","SAME PATTERN, TWO ECOSYSTEMS",
         "OPENAI AGENTS SDK",[("Package: openai-agents",0),("Agent(name, instructions, tools)",1),
          ("Runner.run_sync(agent, prompt)",1),("Delegation via handoffs=[...]",1),
          ("Typed output via output_type=Model",1),("Tool via @function_tool",1)],
         "GOOGLE GEMINI ADK",[("Package: google-adk",0),("Agent(name, model, description, instruction)",1),
          ("Runner + InMemorySessionService",1),("Delegation via sub_agents=[...]",1),
          ("Routing driven by each agent's description",1),("Tool via a plain Python function",1)])
    if t["num"]==5:
        flow_h("How MCP Connects an Agent to Your Tools",
               ["You write an MCP server exposing typed tools",
                "The server runs over stdio (or HTTP)",
                "You register it with an MCP client",
                "The client discovers the tool list",
                "The agent calls tools as it reasons"],
               kicker="WRITE ONCE · USE FROM ANY MCP CLIENT",color=VIOLET)
        architecture_diagram("Hierarchical Sub-Agent Architecture","TOPIC 05 · THREE TIERS","ORCHESTRATOR  (decompose and delegate)",
         [("Sub-Agent 1","Own context and tools"),("Sub-Agent 2","Own context and tools"),("Sub-Agent 3","Own context and tools")],
         tools_label="YOUR MCP SERVER  ·  TYPED TOOLS SHARED BY EVERY SUB-AGENT",accent=VIOLET)
        tile_grid("Why Context Isolation Matters",[
         ("The problem","One agent doing everything fills its context with irrelevant working detail."),
         ("The fix","A sub-agent works in its own context and returns only the conclusion."),
         ("The parent stays clean","The orchestrator sees answers, not transcripts."),
         ("Parallelism","Independent sub-agents run concurrently, cutting wall-clock time."),
         ("Failure containment","One failing sub-agent degrades the result instead of killing the run."),
         ("Testability","Each sub-agent has a narrow contract you can test on its own.")],
         kicker="THE ARCHITECTURAL PAYOFF",cols=2,size=13,accent=VIOLET)

    acts=TOPIC_ACTS[t["num"]]
    # card summary of the labs in this topic
    third=max(1,(len(acts)+2)//3)
    groups=[acts[i:i+third] for i in range(0,len(acts),third)][:3]
    while len(groups)<3: groups.append([])
    cards=[]
    for gi,g in enumerate(groups):
        cards.append((CARD_COLORS[gi],
                      (f"Lab {g[0]['num']}" if len(g)==1 else f"Labs {g[0]['num']}–{g[-1]['num']}") if g else "—",
                      [a["title"] for a in g] if g else ["—"]))
    cards3(f"Hands-On Labs — {t['title']}", cards, kicker="WHAT YOU'LL DO")
    # per activity — overview + workflow diagram + verification (NO step-by-step slides)
    for a in acts:
        activity_overview(f"LAB {a['num']}", a["title"], a["desc"], a["build"], a["services"],
                          kicker=f"TOPIC {t['code']} · HANDS-ON",
                          lab_path=C.lab_slug(a["num"], a["title"]))
        lab_workflow(a)
        test_slide(a["title"], a["test"], kicker=f"LAB {a['num']} · VERIFY")
    # topic recap
    content(f"Recap — {t['title']}",
            ["You can now: "+a["objective"] for a in {x["objective"]:x for x in acts}.values()][:6],
            kicker="TOPIC RECAP", size=17)

# ---------------- CLOSE ----------------
section("WRAP-UP","Course Summary & Next Steps","")
tile_grid("What You Achieved",[
 ("LO1 · Agent Foundations","Built a tool-calling agent with memory and a reusable skill."),
 ("LO2 · Vibe Coding","Applied specify-scaffold-generate-verify-refine with engineered context."),
 ("LO3 · OpenAI Agents SDK","Built and deployed a supervisor-routed multi-agent system."),
 ("LO4 · Gemini Agent SDK","Built and deployed the same pattern on Google's ADK."),
 ("LO5 · MCP & Sub-Agents","Wrote an MCP server and a hierarchical sub-agent architecture.")],
 kicker="LEARNING OUTCOMES",cols=1,size=14)
tile_grid("Where to Go Next",[
 ("Rebuild from memory","Redo the labs without the guide until the patterns are automatic."),
 ("Ship something small","Pick one real task at work and build a two-agent system for it."),
 ("Publish an MCP server","Wrap a capability your team already uses and share it."),
 ("Add evaluation","Write test cases for your agents and measure changes against them."),
 ("Watch the ecosystem","Both SDKs move quickly — follow their changelogs."),
 ("Keep the human gate","Review agent output before it reaches production.")],
 kicker="NEXT STEPS",cols=2,size=14)
tile_grid("Summary",[
 ("From one agent to many","We moved from a single tool-calling agent to hierarchical multi-agent systems."),
 ("Two ecosystems","We built the same architecture twice — OpenAI Agents SDK and Google Gemini ADK."),
 ("Vibe coding as method","Engineered context and a written spec — a build method, not a shortcut."),
 ("MCP as leverage","A capability written once becomes reusable by every agent you write."),
 ("Context isolation","Sub-agents return conclusions, keeping the parent reliable at scale."),
 ("Questions?","Ask anything before we move to the assessment.")],
 kicker="WRAP-UP  ·  Q&A",cols=2,size=14)
cards3("Assessment",[
 (BLUE,"Written (WA)",["Short-Answer Questions (SAQ)","50 minutes","Answer in your own words"]),
 (TEAL,"Practical (PP)",["Multi-agent build tasks","75 minutes","Based on the labs you did"]),
 (VIOLET,"Remember",["Open book — slides and Learner Guide",
                      "Take the Assessment digital attendance (TRAQOM)",
                      "Submit on the LMS: lms-tms.tertiaryinfotech.com"])],kicker="FINAL ASSESSMENT")
flow_h("Assessment Flow",[
 "TRAQOM survey — scan the QR code on the LMS",
 "Assessment digital attendance — scan the SSG QR",
 "Sit WA (SAQ) then PP — open book",
 "Submit your answers on the LMS",
 "Sign the Assessment Summary Record"],kicker="ON ASSESSMENT DAY")
flow_h("Digital Attendance (Mandatory)",[
 "The trainer displays the SSG digital attendance QR code",
 "Scan the QR code with your mobile phone camera",
 "Submit your Assessment digital attendance",
 "Complete the TRAQOM survey on the LMS",
 "Collect your Certificate on the LMS — both are mandatory"],
 kicker="TRAQOM · SSG DIGITAL ATTENDANCE",color=AMBER)
big_statement("Thank You!","You can now design, build and deploy collaborative multi-agent AI systems — with vibe coding as your build method.","HAPPY BUILDING",color=TEAL)

OUT=os.path.join(REPO,"courseware",f"{C.SHORT_TITLE}-{C.VERSION}.pptx")
prs.save(OUT)
print(f"Saved {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
